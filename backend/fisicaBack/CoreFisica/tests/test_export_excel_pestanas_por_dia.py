"""El Excel del Reporte de Asistencia trae DOS pestañas del día seleccionado: DIURNO y
NOCTURNO (volvió al formato anterior; ya no una pestaña por cada día del mes)."""
import io
import json

from django.test import TestCase
from django.contrib.auth.models import User
from openpyxl import load_workbook


def _login(c, u, p):
    return c.post('/api/login/', data=json.dumps({'username': u, 'password': p}),
                  content_type='application/json').json().get('access')


class ExcelDosPestanasTests(TestCase):
    def setUp(self):
        User.objects.create_superuser(username='xls_user', email='e@e.com', password='XlsPass123!')
        self.access = _login(self.client, 'xls_user', 'XlsPass123!')

    def test_solo_diurno_y_nocturno_del_dia(self):
        r = self.client.get('/api/reporte-asistencia/exportar-excel/?fecha=2026-09-03',
                            HTTP_AUTHORIZATION=f'Bearer {self.access}')
        self.assertEqual(r.status_code, 200)
        wb = load_workbook(io.BytesIO(r.content))
        self.assertEqual(wb.sheetnames, ['DIURNO', 'NOCTURNO'])
        self.assertIn('ASISTENCIA GENERAL 03-09-2026.xlsx', r['Content-Disposition'])

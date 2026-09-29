"""El descargable de asignaciones NO debe duplicar los SACAFRANCO.

Bug: el export traía las filas de sacafranco con mes<=actual, así que sumaba las de meses
anteriores y la misma persona salía DOS veces (una con datos y otra vacía). El endpoint en
vivo ya filtraba por mes exacto; este test asegura que el descargable haga lo mismo.
"""
import io
import json

from django.test import TestCase
from django.contrib.auth.models import User
from openpyxl import load_workbook

from CoreFisica.models import Persona, SacafrancoFila


def _login(c, u, p):
    return c.post('/api/login/', data=json.dumps({'username': u, 'password': p}),
                  content_type='application/json').json().get('access')


class ExportSacafrancoNoDuplicaTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_superuser(
            username='exp_user', email='e@e.com', password='ExpPass123!'
        )
        self.access = _login(self.client, 'exp_user', 'ExpPass123!')
        self.mes, self.anio = 7, 2026
        self.persona = Persona.objects.create(
            nombres='MAURICIO EFRAIN', apellidos='BRAVO CABRERA',
            cedula='0920388246', tipo='SACAFRANCO',
        )
        # Fila del MES ACTUAL (con datos) + una del MES ANTERIOR (la que duplicaba).
        SacafrancoFila.objects.create(persona=self.persona, mes=self.mes, anio=self.anio, orden=1)
        SacafrancoFila.objects.create(persona=self.persona, mes=self.mes - 1, anio=self.anio, orden=1)

    def _contar_en_hoja(self, ws, cedula):
        # La cédula va en la columna 6 (F). Cuenta cuántas filas la tienen.
        return sum(
            1 for row in ws.iter_rows(min_col=6, max_col=6)
            if str(row[0].value or '').strip() == cedula
        )

    def test_descargable_no_duplica_sacafranco(self):
        r = self.client.get(
            f'/api/reporte-asignaciones/?mes={self.mes}&anio={self.anio}',
            HTTP_AUTHORIZATION=f'Bearer {self.access}'
        )
        self.assertEqual(r.status_code, 200)
        self.assertIn('spreadsheet', r['Content-Type'])

        wb = load_workbook(io.BytesIO(r.content))
        self.assertIn('Asignaciones y Calendario', wb.sheetnames)
        ws = wb['Asignaciones y Calendario']

        cont = self._contar_en_hoja(ws, self.persona.cedula)
        self.assertEqual(
            cont, 1,
            f'El SACAFRANCO no debe duplicarse en el descargable (aparece {cont} veces)'
        )

    def test_solo_trae_la_fila_del_mes_exacto(self):
        # Aunque exista fila de meses anteriores, en la BD hay 2 filas pero el descargable
        # (mes exacto) solo debe reflejar la del mes consultado -> 1 fila del sacafranco.
        self.assertEqual(SacafrancoFila.objects.filter(persona=self.persona).count(), 2)
        r = self.client.get(
            f'/api/reporte-asignaciones/?mes={self.mes}&anio={self.anio}',
            HTTP_AUTHORIZATION=f'Bearer {self.access}'
        )
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb['Asignaciones y Calendario']
        self.assertEqual(self._contar_en_hoja(ws, self.persona.cedula), 1)

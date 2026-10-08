"""Servicios Adicionales: lista de los ADICIONALES de un rango de fechas, sin valores.

Sale de la sección ADICIONALES del Reporte de Guardia, que se llena sola desde la asistencia (estado ADICIONAL
con su reemplazo); también trae las filas que se agregaron o corrigieron a mano en el Reporte de Guardia. Las
columnas son las mismas de esa sección: Cliente, Puesto, Nombre y Proviene, más la Fecha y el Turno.
Solo lectura: para corregir un registro se edita la asistencia o el Reporte de Guardia.
"""
import datetime
import unicodedata
from io import BytesIO

from django.http import HttpResponse
from rest_framework import status
from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response

from ..models import ReporteGuardia

PERMISO = 'CoreFisica.view_servicioadicional'   # permiso propio de solo lectura
MAX_DIAS = 366   # rango máximo que se puede pedir de una vez


def _fecha(v):
    try:
        return datetime.date.fromisoformat(str(v)[:10]) if v else None
    except ValueError:
        return None


def _norm(s):
    s = unicodedata.normalize('NFD', str(s or '').lower())
    return ''.join(c for c in s if unicodedata.category(c) != 'Mn')


def _filas(request):
    """(filas, error). Filtros: ?desde=&hasta= (o ?fecha=), ?turno=Diurno|Nocturno (sin turno = ambos), ?q=."""
    fecha = _fecha(request.GET.get('fecha'))
    desde = _fecha(request.GET.get('desde')) or fecha
    hasta = _fecha(request.GET.get('hasta')) or fecha or desde
    if not desde or not hasta:
        return None, 'Indica la fecha o el rango (desde / hasta).'
    if hasta < desde:
        desde, hasta = hasta, desde
    if (hasta - desde).days > MAX_DIAS:
        return None, f'El rango no puede pasar de {MAX_DIAS} días.'
    qs = ReporteGuardia.objects.filter(seccion='ADICIONALES', fecha__gte=desde, fecha__lte=hasta)
    turno = request.GET.get('turno')
    if turno in ('Diurno', 'Nocturno'):
        qs = qs.filter(turno=turno)
    filas = [{
        'id': r.id,
        'fecha': r.fecha.isoformat(),
        'turno': r.turno or '',
        'cliente': r.cliente or '',
        'puesto': r.puesto or '',
        'persona_nombre': r.persona_nombre or '',
        'proviene': r.proviene or '',
    } for r in qs.order_by('fecha', 'turno', 'cliente', 'puesto', 'id')]
    tokens = _norm(request.GET.get('q')).split()
    if tokens:
        def _ok(f):
            txt = _norm(' '.join(f[k] for k in ('cliente', 'puesto', 'persona_nombre', 'proviene')))
            return all(t in txt for t in tokens)
        filas = [f for f in filas if _ok(f)]
    return filas, None


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def listar_servicios_adicionales(request):
    if not request.user.has_perm(PERMISO):
        return Response({'error': 'No autorizado'}, status=status.HTTP_403_FORBIDDEN)
    filas, error = _filas(request)
    if error:
        return Response({'error': error}, status=status.HTTP_400_BAD_REQUEST)
    return Response(filas)


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def exportar_servicios_adicionales_excel(request):
    """Excel con las mismas columnas de la tabla (fechas, turno y búsqueda aplicados)."""
    if not request.user.has_perm(PERMISO):
        return Response({'error': 'No autorizado'}, status=status.HTTP_403_FORBIDDEN)
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    filas, error = _filas(request)
    if error:
        return Response({'error': error}, status=status.HTTP_400_BAD_REQUEST)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'ADICIONALES'
    columnas = [('Nº', 6), ('FECHA', 12), ('TURNO', 11), ('CLIENTE', 34), ('PUESTO', 34),
                ('NOMBRES Y APELLIDOS', 38), ('PROVIENE', 30)]
    borde = Border(*(Side(style='thin', color='999999'),) * 4)
    for c, (titulo, ancho) in enumerate(columnas, start=1):
        cell = ws.cell(1, c, titulo)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='1F4E78')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = borde
        ws.column_dimensions[get_column_letter(c)].width = ancho
    for n, f in enumerate(filas, start=1):
        fecha = datetime.date.fromisoformat(f['fecha']).strftime('%d/%m/%Y')
        valores = [n, fecha, f['turno'], f['cliente'], f['puesto'], f['persona_nombre'], f['proviene']]
        for c, v in enumerate(valores, start=1):
            cell = ws.cell(n + 1, c, v)
            cell.border = borde
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.freeze_panes = 'A2'
    out = BytesIO()
    wb.save(out)
    resp = HttpResponse(out.getvalue(),
                        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    resp['Content-Disposition'] = 'attachment; filename="servicios_adicionales.xlsx"'
    return resp

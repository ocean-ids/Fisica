import datetime
from io import BytesIO
from django.db.models import Q
from django.http import HttpResponse, JsonResponse
from rest_framework.decorators import api_view, permission_classes
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework import status
from django.shortcuts import get_object_or_404
import openpyxl
from openpyxl.styles import Alignment, Border, Side, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils.units import pixels_to_EMU
from ..models import ReporteGuardia, Asignacion
from ..serializers import ReporteGuardiaSerializer

TURNOS = ('Diurno', 'Nocturno')

# Encabezado del formato FR REPORTE DE GUARDIA (para el descargable).
FR_TITULO = 'FR REPORTE DE GUARDIA'
FR_VERSION = '.03'
FR_FECHA_APROBACION = '2021-12-14'
# Anchos de columna A..H. La A es angosta porque la etiqueta de seccion va rotada (vertical).
FR_ANCHOS = {1: 5.0, 2: 5.3, 3: 50.3, 4: 35.1, 5: 56.3, 6: 52.3, 7: 28.0, 8: 21.4}

# Secciones del FR, en orden. Cada columna extra: (encabezado, campo_modelo, col_ini, col_fin).
# La última columna de cada sección se combina hasta H, igual que el formato original.
FR_SECCIONES = [
    ('DOBLADAS',     'DOBLADAS',     [('1 NOMBRE Y 2 APELLIDOS', 'persona_nombre', 5, 5), ('PROVIENE', 'proviene', 6, 6), ('VALOR', 'valor', 7, 8)]),
    ('ADICIONALES',  'ADICIONALES',  [('1 NOMBRE Y 2 APELLIDOS', 'persona_nombre', 5, 5), ('PROVIENE', 'proviene', 6, 8)]),
    ('ADELANTOS',    'ADELANTOS',    [('1 NOMBRE Y 2 APELLIDOS', 'persona_nombre', 5, 5), ('PROVIENE', 'proviene', 6, 6), ('TIPO', 'tipo', 7, 8)]),
    ('NO_CUBIERTOS', 'NO CUBIERTOS', [('AUTORIZACION', 'autorizacion', 5, 5), ('MOTIVO', 'motivo', 6, 8)]),
    ('FALTOS',       'FALTOS',       [('1 NOMBRE Y 2 APELLIDOS', 'persona_nombre', 5, 5), ('MOTIVO', 'motivo', 6, 8)]),
    ('HUECA',        'HUECA',        [('MOTIVO', 'motivo', 5, 7), ('FECHA', 'fecha_evento', 8, 8)]),
    ('APOYO',        'APOYO',        [('1 NOMBRE Y 2 APELLIDOS', 'persona_nombre', 5, 5), ('PROVIENE', 'proviene', 6, 6), ('MOTIVO', 'motivo', 7, 8)]),
]

# Campos de contenido que el usuario puede editar a mano en el reporte de guardia.
EDITABLE_CONTENT_FIELDS = (
    'cliente', 'puesto', 'persona_nombre', 'proviene', 'valor', 'tipo',
    'autorizacion', 'motivo', 'fecha_evento',
)


def _sync_no_cubiertos(fecha, turno):
    """Refleja en NO_CUBIERTOS los puestos (asignaciones) SIN persona para esa
    fecha/turno. Conserva la autorizacion/motivo escritos a mano. Idempotente:
    crea/actualiza una fila auto por asignacion vacante y borra las que ya no aplican."""
    try:
        fecha_obj = fecha if isinstance(fecha, datetime.date) else datetime.date.fromisoformat(str(fecha))
    except (TypeError, ValueError):
        return
    if turno not in TURNOS:
        return

    from .reporte_asistencia_views import _calendar_dnf_for_date

    mes, anio = fecha_obj.month, fecha_obj.year
    month_start = fecha_obj.replace(day=1)
    if mes == 12:
        month_end = datetime.date(anio + 1, 1, 1) - datetime.timedelta(days=1)
    else:
        month_end = datetime.date(anio, mes + 1, 1) - datetime.timedelta(days=1)

    # Vacantes: asignaciones ACTIVAS sin persona, puesto abierto, que aplican al mes.
    vac = Asignacion.objects.select_related('cliente', 'puesto').prefetch_related('puesto__horarios').filter(
        estado='ACTIVO', persona__isnull=True, puesto__activo=True
    ).filter(
        Q(mes=mes, anio=anio) |
        (Q(recurring=True) & Q(start_date__lte=month_end) & (Q(end_date__isnull=True) | Q(end_date__gte=month_start)))
    ).exclude(
        # Historial por día: una vacante con fecha de alta no aparece en días anteriores
        # a su creación. NULL = sin corte (vigente todo el mes, como antes).
        vigente_desde__isnull=False, vigente_desde__gt=fecha_obj
    )

    dnf = _calendar_dnf_for_date(fecha_obj)      # {asignacion_id: 'D'|'N'|'F'}
    turno_letter = 'D' if turno == 'Diurno' else 'N'
    dia = fecha_obj.weekday() + 1                 # 1=Lunes ... 7=Domingo

    def _aplica(a):
        letra = dnf.get(a.id)
        if letra in ('D', 'N'):
            return letra == turno_letter          # el calendario manda si existe
        if letra == 'F':
            return False
        # Sin calendario: derivar el turno del horario del puesto ese dia.
        turnos = set()
        horarios = a.puesto.horarios.all() if a.puesto else []
        for h in horarios:
            if getattr(h, 'dia', None) != dia:
                continue
            t = (getattr(h, 'turno', '') or '').strip().lower()
            if t.startswith('d'):
                turnos.add('Diurno')
            elif t.startswith('n'):
                turnos.add('Nocturno')
            else:
                turnos.update({'Diurno', 'Nocturno'})   # 24h / ambos
        return turno in turnos

    vigentes = [a for a in vac if _aplica(a)]
    vigentes_ids = {a.id for a in vigentes}

    # Borrar filas auto que ya no aplican (cubiertas, cerradas o fuera de turno).
    ReporteGuardia.objects.filter(
        fecha=fecha_obj, turno=turno, seccion='NO_CUBIERTOS', auto=True
    ).exclude(asignacion_ref_id__in=vigentes_ids).delete()

    # Upsert por asignacion (conserva autorizacion/motivo escritos a mano).
    existentes = {
        r.asignacion_ref_id: r
        for r in ReporteGuardia.objects.filter(
            fecha=fecha_obj, turno=turno, seccion='NO_CUBIERTOS', auto=True
        )
    }
    for a in vigentes:
        cliente = getattr(a.cliente, 'nombre_comercial', '') or ''
        puesto = getattr(a.puesto, 'nombre', '') or ''
        row = existentes.get(a.id)
        if row:
            ov = row.overrides or {}
            cambios = []
            # No pisar los campos que el usuario editó a mano.
            if 'cliente' not in ov and row.cliente != cliente:
                row.cliente = cliente
                cambios.append('cliente')
            if 'puesto' not in ov and row.puesto != puesto:
                row.puesto = puesto
                cambios.append('puesto')
            if cambios:
                row.save(update_fields=cambios)
        else:
            ReporteGuardia.objects.create(
                fecha=fecha_obj, turno=turno, seccion='NO_CUBIERTOS',
                cliente=cliente, puesto=puesto, auto=True, asignacion_ref=a,
            )


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def listar_reporte_guardia(request):
    """Filas del reporte por ?fecha=YYYY-MM-DD&turno=Diurno|Nocturno (opcional ?seccion=)."""
    fecha = request.GET.get('fecha')
    turno = request.GET.get('turno')
    seccion = request.GET.get('seccion')
    # NO_CUBIERTOS ahora es MANUAL: ya no se extrae automáticamente de las asignaciones
    # vacantes. Se crea/edita/elimina a mano desde el Reporte de Guardia.
    qs = ReporteGuardia.objects.select_related('persona_ref')
    if fecha:
        qs = qs.filter(fecha=fecha)
    if turno in TURNOS:
        qs = qs.filter(turno=turno)
    if seccion:
        qs = qs.filter(seccion=seccion)
    return Response(ReporteGuardiaSerializer(qs, many=True).data)

@api_view(['POST'])
@permission_classes([IsAuthenticated])
def crear_reporte_guardia(request):
    """Crea una fila manual del reporte (p. ej. APOYO). No se sincroniza con nada;
    solo existe en el reporte de guardia."""
    s = ReporteGuardiaSerializer(data=request.data)
    s.is_valid(raise_exception=True)
    s.save(auto=False)   # manual: nunca lo toca la sincronizacion automatica
    return Response(s.data, status=status.HTTP_201_CREATED)


@api_view(['PUT', 'PATCH'])
@permission_classes([IsAuthenticated])
def actualizar_reporte_guardia(request, id):
    fila = get_object_or_404(ReporteGuardia, id=id)
    s = ReporteGuardiaSerializer(fila, data=request.data, partial=True)
    s.is_valid(raise_exception=True)
    fila = s.save()
    # En filas auto (generadas desde asistencia), registrar los campos editados a
    # mano para conservarlos cuando la fila se regenere. Las manuales (APOYO) no se
    # regeneran, así que no hace falta.
    if fila.auto:
        ov = dict(fila.overrides or {})
        for campo in request.data:
            if campo in EDITABLE_CONTENT_FIELDS:
                ov[campo] = request.data.get(campo)
        if ov != (fila.overrides or {}):
            fila.overrides = ov
            fila.save(update_fields=['overrides'])
    return Response(ReporteGuardiaSerializer(fila).data)


@api_view(['DELETE'])
@permission_classes([IsAuthenticated])
def eliminar_reporte_guardia(request, id):
    fila = get_object_or_404(ReporteGuardia, id=id)
    fila.delete()
    return Response(status=status.HTTP_204_NO_CONTENT)


@api_view(['POST'])
@permission_classes([IsAuthenticated])
def regenerar_reporte_guardia(request):
    """Regenera BAJO DEMANDA las filas auto del reporte de guardia desde la
    asistencia de una fecha (faltos, dobladas, adicionales, adelantos, huecas).
    Vuelve a traer esas filas desde la asistencia y reaplica las ediciones a mano;
    las filas MANUALES (auto=False) no se tocan. Se dispara con el boton
    'Regenerar desde asistencia' del Reporte de Guardia."""
    fecha = request.data.get('fecha') or request.GET.get('fecha')
    try:
        fecha_obj = fecha if isinstance(fecha, datetime.date) else datetime.date.fromisoformat(str(fecha))
    except (TypeError, ValueError):
        return Response({'error': 'fecha invalida'}, status=status.HTTP_400_BAD_REQUEST)

    from .reporte_asistencia_views import (
        _sync_reporte_guardia, _sync_reporte_guardia_sacafranco,
        _sync_hueca_reporte_guardia_sacafranco, _sync_frtrabajado_dobladas,
    )
    from ..models import ReporteAsistencia, SacafrancoAsistencia

    # Se incluyen TODAS las asignaciones con datos guardados ese día (tengan estado ACTIVO
    # o INACTIVO), igual que el Reporte de Asistencia: un puesto cerrado (INACTIVO) con
    # FALTO/hueca/FR-TRABAJADO ese día debe reflejarse en el Reporte de Guardia. Antes solo
    # se procesaban las ACTIVO, por eso esos registros no generaban FALTOS/DOBLADAS/HUECA.
    overrides = ReporteAsistencia.objects.select_related('asignacion').filter(
        fecha_reporte=fecha_obj, asignacion__isnull=False
    )
    for ov in overrides:
        if not ov.asignacion:
            continue
        try:
            _sync_reporte_guardia(ov, ov.asignacion, fecha_obj)
        except Exception:
            pass
        # FR/TRABAJADO -> DOBLADAS (fila manual). Se reconstruye aquí para que un registro
        # que no se guardó bien también aparezca al regenerar.
        try:
            _sync_frtrabajado_dobladas(ov, ov.asignacion, fecha_obj)
        except Exception:
            pass

    # Sacafranco: su asistencia esta en SacafrancoAsistencia (no tiene asignacion).
    for sa in SacafrancoAsistencia.objects.select_related('sacafranco_fila', 'reemplazo').filter(fecha=fecha_obj):
        try:
            _sync_reporte_guardia_sacafranco(sa, fecha_obj)
            _sync_hueca_reporte_guardia_sacafranco(sa, fecha_obj)
        except Exception:
            pass
    return Response({'ok': True})



def _fr_header_block(ws, top, fecha_obj, turno, border, logo_path):
    """Bloque de encabezado FR (logo + titulo + version + FECHA/TURNO) de un turno.
    Ocupa las filas top..top+5. Devuelve la fila de la primera seccion."""
    bold16 = Font(bold=True, size=16)
    reg11 = Font(size=11)
    center = Alignment(horizontal='center', vertical='center')
    for r in range(top, top + 4):
        ws.row_dimensions[r].height = 21.6

    # Logo (A:C, 4 filas) — centrado dentro del rectangulo.
    ws.merge_cells(start_row=top, start_column=1, end_row=top + 3, end_column=3)
    if logo_path:
        try:
            w_px, h_px = 150, 54
            cols_px = [round(FR_ANCHOS[c] * 7) + 5 for c in (1, 2, 3)]  # ancho px de A,B,C
            row_px = round(21.6 * 96 / 72)                              # alto px por fila
            left_abs = max(int((sum(cols_px) - w_px) / 2), 0)          # x centrado (desde A)
            top_abs = max(int((row_px * 4 - h_px) / 2), 0)             # y centrado (4 filas)
            # Localizar en que columna/fila cae ese punto y el offset dentro de ella.
            acol, cum = 0, 0
            for i, wpx in enumerate(cols_px):
                if left_abs < cum + wpx:
                    acol, off_x = i, left_abs - cum
                    break
                cum += wpx
            else:
                acol, off_x = 2, 0
            arow = top - 1 + (top_abs // row_px)
            off_y = top_abs % row_px
            img = XLImage(str(logo_path))
            marker = AnchorMarker(col=acol, colOff=pixels_to_EMU(off_x), row=arow, rowOff=pixels_to_EMU(off_y))
            img.anchor = OneCellAnchor(_from=marker, ext=XDRPositiveSize2D(pixels_to_EMU(w_px), pixels_to_EMU(h_px)))
            ws.add_image(img)
        except Exception:
            pass
    # Titulo (D:F)
    ws.merge_cells(start_row=top, start_column=4, end_row=top + 3, end_column=6)
    tc = ws.cell(row=top, column=4, value=FR_TITULO)
    tc.font = bold16
    tc.alignment = center
    # Version / Fecha de aprobacion (G/H)
    ws.merge_cells(start_row=top, start_column=7, end_row=top + 1, end_column=7)
    ws.merge_cells(start_row=top, start_column=8, end_row=top + 1, end_column=8)
    ws.merge_cells(start_row=top + 2, start_column=7, end_row=top + 3, end_column=7)
    ws.merge_cells(start_row=top + 2, start_column=8, end_row=top + 3, end_column=8)
    ws.cell(row=top, column=7, value='Version:').font = reg11
    ws.cell(row=top, column=8, value=FR_VERSION).font = reg11
    ws.cell(row=top + 2, column=7, value='Fecha de aprobacion:').font = reg11
    ws.cell(row=top + 2, column=8, value=FR_FECHA_APROBACION).font = reg11
    for r in range(top, top + 4):
        for c in range(1, 9):
            ws.cell(row=r, column=c).border = border
    for cc in (7, 8):
        ws.cell(row=top, column=cc).alignment = center
        ws.cell(row=top + 2, column=cc).alignment = center

    # Fila FECHA / TURNO (top+5); top+4 queda en blanco
    fr = top + 5
    ws.merge_cells(start_row=fr, start_column=1, end_row=fr, end_column=2)
    ws.merge_cells(start_row=fr, start_column=3, end_row=fr, end_column=5)
    ws.merge_cells(start_row=fr, start_column=6, end_row=fr, end_column=8)
    ws.cell(row=fr, column=1, value='FECHA: ').font = reg11
    dc = ws.cell(row=fr, column=3, value=fecha_obj)
    dc.number_format = 'yyyy-mm-dd'
    ws.cell(row=fr, column=6, value='TURNO: %s' % turno.upper()).font = reg11
    for c in range(1, 9):
        cell = ws.cell(row=fr, column=c)
        cell.border = border
        cell.alignment = center
    return fr + 1


def _fr_write_secciones(ws, start_row, fecha_obj, turno, border):
    """Escribe las 7 secciones del FR desde start_row. Devuelve la fila siguiente."""
    reg11 = Font(size=11)
    bold11 = Font(bold=True, size=11)
    bold12 = Font(bold=True, size=12)
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    r = start_row
    for seccion_key, etiqueta, extras in FR_SECCIONES:
        header_row = r
        # Encabezado de columnas
        ws.cell(row=r, column=2, value='N°')
        ws.cell(row=r, column=3, value='CLIENTE')
        ws.cell(row=r, column=4, value='PUESTO')
        for h, _campo, ci, cf in extras:
            ws.cell(row=r, column=ci, value=h)
            if cf > ci:
                ws.merge_cells(start_row=r, start_column=ci, end_row=r, end_column=cf)
        for c in range(2, 9):
            cell = ws.cell(row=r, column=c)
            cell.font = bold11
            cell.alignment = center
            cell.border = border
        r += 1

        filas = ReporteGuardia.objects.filter(
            fecha=fecha_obj, turno=turno, seccion=seccion_key
        ).order_by('orden', 'id')
        for n, fila in enumerate(filas, start=1):
            ws.cell(row=r, column=2, value=n)
            ws.cell(row=r, column=3, value=fila.cliente or '')
            ws.cell(row=r, column=4, value=fila.puesto or '')
            for _h, campo, ci, cf in extras:
                val = getattr(fila, campo, None)
                if campo == 'valor':
                    val = float(val) if val is not None else 0
                elif campo == 'fecha_evento':
                    pass
                else:
                    val = val or ''
                cell = ws.cell(row=r, column=ci, value=val)
                if campo == 'fecha_evento' and val:
                    cell.number_format = 'yyyy-mm-dd'
                if cf > ci:
                    ws.merge_cells(start_row=r, start_column=ci, end_row=r, end_column=cf)
            for c in range(2, 9):
                cell = ws.cell(row=r, column=c)
                cell.font = reg11
                cell.alignment = center
                cell.border = border
            r += 1

        # Reservar un minimo de filas por seccion para que la etiqueta vertical (rotada)
        # quepa completa aunque haya pocas/ninguna fila (como el formato FR original).
        MIN_FILAS = 6
        data_escritas = r - (header_row + 1)
        for _ in range(max(MIN_FILAS - data_escritas, 0)):
            for c in range(2, 9):
                ws.cell(row=r, column=c).border = border
            r += 1

        last_row = r - 1
        # Etiqueta de seccion en A, combinada verticalmente (header + filas de datos).
        if last_row > header_row:
            ws.merge_cells(start_row=header_row, start_column=1, end_row=last_row, end_column=1)
        acell = ws.cell(row=header_row, column=1, value=etiqueta)
        acell.font = bold12
        # Rotada 90° (vertical), como el formato modelo; no se corta.
        acell.alignment = Alignment(horizontal='center', vertical='center', text_rotation=90, wrap_text=False)
        for rr in range(header_row, last_row + 1):
            ws.cell(row=rr, column=1).border = border
    return r


@api_view(['GET'])
@permission_classes([IsAuthenticated])
def exportar_reporte_guardia_excel(request):
    """Descargable del Reporte de Guardia en formato FR (una hoja por dia, turnos
    Diurno y Nocturno, con las 7 secciones). Parametro: ?fecha=YYYY-MM-DD."""
    fecha = request.GET.get('fecha')
    try:
        fecha_obj = datetime.date.fromisoformat(str(fecha))
    except (TypeError, ValueError):
        return JsonResponse({'error': 'fecha invalida'}, status=status.HTTP_400_BAD_REQUEST)

    from .reporte_asistencia_views import _find_logo_path
    logo_path = _find_logo_path()

    thin = Side(border_style='thin', color='000000')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    bold = Font(bold=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = fecha_obj.strftime('%d-%m-%Y')
    for c, w in FR_ANCHOS.items():
        ws.column_dimensions[get_column_letter(c)].width = w

    def firmas(row):
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        ws.merge_cells(start_row=row, start_column=4, end_row=row, end_column=6)
        ws.merge_cells(start_row=row, start_column=7, end_row=row, end_column=8)
        ws.cell(row=row, column=1, value='ELABORA:').font = bold
        ws.cell(row=row, column=4, value='REVISA:').font = bold
        ws.cell(row=row, column=7, value='AUTORIZA:').font = bold

    # TURNO DIURNO
    r = _fr_header_block(ws, 1, fecha_obj, 'Diurno', border, logo_path)
    r = _fr_write_secciones(ws, r, fecha_obj, 'Diurno', border)
    firmas(r + 1)

    # TURNO NOCTURNO (nuevo bloque de encabezado con su logo)
    top2 = r + 4
    r = _fr_header_block(ws, top2, fecha_obj, 'Nocturno', border, logo_path)
    r = _fr_write_secciones(ws, r, fecha_obj, 'Nocturno', border)
    firmas(r + 1)

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="reporte_guardia_%s.xlsx"' % fecha_obj.isoformat()
    out = BytesIO()
    wb.save(out)
    out.seek(0)
    response.write(out.getvalue())
    return response
